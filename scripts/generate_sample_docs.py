"""Generate the PDF and DOCX sample documents (dev-only: needs reportlab + python-docx)."""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "data" / "documents"

SECURITY = [
    ("Password requirements", "Passwords must be at least 14 characters and are rotated only on suspected compromise. "
     "Multi-factor authentication is mandatory for email, VPN and the admin console."),
    ("Data classification", "Data is classified as Public, Internal, Confidential or Restricted. Customer personal data and "
     "financial records are Restricted and may only be shared with approval from the Security Officer."),
    ("Incident reporting", "Any suspected security incident must be reported to security@acme.example within 1 hour of discovery. "
     "Do not attempt to investigate or delete evidence yourself."),
    ("Acceptable use of AI tools", "Employees must not paste Confidential or Restricted data into external AI services. "
     "Only approved, locally hosted AI tools may process company data."),
]

doc = Document()
doc.add_heading("Acme Corp Information Security Policy", 0)
for title, body in SECURITY:
    doc.add_heading(title, 1)
    doc.add_paragraph(body)
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Class", "Retention"
for cls, ret in [("Internal", "3 years"), ("Confidential", "7 years"), ("Restricted", "10 years")]:
    row = table.add_row().cells
    row[0].text, row[1].text = cls, ret
doc.save(OUT / "security_policy.docx")

styles = getSampleStyleSheet()
story = [Paragraph("Acme Corp Product Warranty and Support", styles["Title"])]
sections = [
    ("Warranty coverage", "All Acme computers carry a 24-month limited warranty covering manufacturing defects. "
     "Accessories, audio devices and wearables carry a 12-month warranty. Accidental damage is not covered unless the customer "
     "purchased the Care+ plan."),
    ("Care+ plan", "Care+ costs 12% of the product price per year and covers one accidental-damage repair per year, "
     "with on-site service for enterprise customers."),
    ("Support hours", "Support is available Monday to Friday, 08:00-20:00 UTC, by email and chat. Enterprise customers "
     "have a 24/7 phone line with a 1-hour response target for critical issues."),
    ("Service levels", "Standard tickets receive a first response within 1 business day. Critical outages for enterprise "
     "customers receive a response within 1 hour."),
]
for title, body in sections:
    story += [Paragraph(title, styles["Heading2"]), Paragraph(body, styles["BodyText"]), Spacer(1, 12)]
SimpleDocTemplate(str(OUT / "warranty_and_support.pdf"), pagesize=A4).build(story)
print("written to", OUT)
