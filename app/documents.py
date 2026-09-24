"""Loading and chunking of PDF / DOCX / TXT / Markdown files."""
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


@dataclass
class Section:
    text: str
    page: int | None = None


# Bump when chunking/metadata changes so already-indexed documents are re-embedded on startup.
CHUNKER_VERSION = 2


@dataclass
class Chunk:
    text: str
    page: int | None
    index: int
    section: str | None = None       # nearest level-1/2 heading, e.g. "Database"
    heading: str | None = None       # closest heading, e.g. "Axios Instance"
    heading_path: str | None = None  # "Frontend Standards > API & Services > Axios Instance"

    def embedding_text(self, filename: str) -> str:
        """What gets embedded: body plus its location, so "Prisma 7" is anchored to "Database"."""
        where = self.heading_path or filename
        return f"{filename} > {where}\n{self.text}" if self.heading_path else f"{filename}\n{self.text}"


def doc_id_for(filename: str) -> str:
    """Stable id derived from the filename, so re-ingesting a file replaces its old chunks."""
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"doc:{filename}"))


def load_sections(path: Path) -> list[Section]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or '(none)'}")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [Section(p.extract_text() or "", i + 1) for i, p in enumerate(reader.pages)]
    if ext == ".docx":
        from docx import Document

        return [Section(_docx_to_markdown(Document(str(path))))]
    return [Section(path.read_text(encoding="utf-8", errors="replace"))]


_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")
_FENCE = re.compile(r"^\s*(```|~~~)")
_NUMBERING = re.compile(r"^\d+(?:\.\d+)*[.)]?\s+")
# Navigation-only sections would match almost every keyword query, so they are not indexed.
_SKIP_HEADINGS = {"table of contents", "contents", "toc"}


@dataclass
class MarkdownSection:
    level: int                      # 0 = text before the first heading
    heading: str | None
    path: list[tuple[int, str]]     # [(level, title), ...] from the document root
    body: str


def _clean_title(raw: str) -> str:
    return _NUMBERING.sub("", raw.strip().strip("`*_ ")) or raw.strip()


def split_markdown(text: str) -> list[MarkdownSection]:
    """Split on ATX headings (# .. ######), ignoring '#' lines inside fenced code blocks.

    A section is a heading plus the text up to the next heading of any level; its path records the
    parent headings. Headings with no text of their own (pure containers) yield no section.
    """
    sections: list[MarkdownSection] = []
    stack: list[tuple[int, str]] = []
    cur = MarkdownSection(0, None, [], "")
    lines: list[str] = []
    in_fence = False

    def flush():
        cur.body = "\n".join(lines).strip()
        if cur.body and (cur.heading or "").lower() not in _SKIP_HEADINGS:
            sections.append(cur)

    for line in text.splitlines():
        if _FENCE.match(line):
            in_fence = not in_fence
        m = None if in_fence else _HEADING.match(line)
        if m:
            flush()
            level, title = len(m.group(1)), _clean_title(m.group(2))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            cur, lines = MarkdownSection(level, title, list(stack), ""), []
        else:
            lines.append(line)
    flush()
    return sections


def _section_name(sec: MarkdownSection) -> str | None:
    """Nearest ancestor (or self) that is a level-1/2 heading, else the heading itself."""
    for level, title in reversed(sec.path):
        if level <= 2:
            return title
    return sec.heading


_HEADED_FORMATS = {".md", ".markdown", ".docx"}  # docx headings are converted to Markdown before this point


# DOCX paragraph style -> Markdown heading level, so any headed .docx reuses the Markdown section
# splitter below (unifies the two formats instead of duplicating the section-assembly logic).
_DOCX_HEADING_STYLE = re.compile(r"^(Title|Heading (\d+))$", re.I)


def _docx_heading_level(style_name: str | None) -> int | None:
    m = _DOCX_HEADING_STYLE.match(style_name or "")
    if not m:
        return None
    return 1 if m.group(1).lower() == "title" else int(m.group(2)) + 1


def _docx_to_markdown(doc) -> str:
    """Represent a .docx as Markdown text: heading-styled paragraphs become '#' lines (any level, any
    wording), everything else stays plain text, and table rows become bullet lines. A document with no
    heading styles at all becomes plain text with no '#' lines, which chunk_sections then treats as a
    single un-headed section (the generic fallback)."""
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        level = _docx_heading_level(p.style.name if p.style else None)
        lines.append(f"{'#' * min(level, 6)} {text}" if level else text)
    for table in doc.tables:
        for row in table.rows:
            cells = " | ".join(c.text.strip() for c in row.cells)
            if cells.strip("| "):
                lines.append(f"- {cells}")
    return "\n\n".join(lines)


def chunk_sections(sections: list[Section], suffix: str = ".txt") -> list[Chunk]:
    """Markdown and DOCX (already converted to Markdown, see _docx_to_markdown): one chunk per heading
    section, split further only if too large, keeping its heading metadata. A document with no headings at
    all naturally falls back to one un-headed section, chunked the same way as PDF/TXT below.
    Other formats: recursive character splitting per page/section."""
    splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", ". ", " ", ""],
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    chunks: list[Chunk] = []
    for section in sections:
        if suffix in _HEADED_FORMATS:
            for md in split_markdown(section.text):
                path = " > ".join(t for _, t in md.path) or None
                for piece in splitter.split_text(md.body):
                    if piece.strip():
                        chunks.append(Chunk(piece.strip(), section.page, len(chunks),
                                            _section_name(md), md.heading, path))
        else:
            for piece in splitter.split_text(section.text):
                if piece.strip():
                    chunks.append(Chunk(piece.strip(), section.page, len(chunks)))
    return chunks
